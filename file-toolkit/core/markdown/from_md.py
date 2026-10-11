"""Markdown 转 Word（.docx）/ 网页（.html）/ PDF。

用 markdown-it-py 解析成语法树（CommonMark + 表格 + 删除线 + 任务列表），再分别生成。
本地图片按 MD 文件所在目录解析：Word / PDF 直接嵌入，HTML 以 base64 内嵌成单文件。
"""
from __future__ import annotations

import base64
import html as html_lib
import io
import mimetypes
import re
import time
from pathlib import Path
from urllib.parse import unquote, urlparse

from core.markdown._common import read_text_file
from core.models import ProgressCallback, TaskResult, TaskStatus
from core.paths import unique_path
from core.task_control import TaskCancelled, check_cancelled

TARGETS = ("docx", "html", "pdf")
_TASK_RE = re.compile(r"^\[([ xX])\]\s+")
_MONO = "Consolas"


def _parser():
    from markdown_it import MarkdownIt

    return MarkdownIt("commonmark", {"html": True}).enable(["table", "strikethrough"])


def _tree(text: str):
    from markdown_it.tree import SyntaxTreeNode

    return SyntaxTreeNode(_parser().parse(text))


def _local_image(src: str, base_dir: Path) -> Path | None:
    parsed = urlparse(src)
    if parsed.scheme and len(parsed.scheme) > 1 and parsed.scheme != "file":
        return None
    path = Path(unquote(parsed.path if parsed.scheme == "file" else src))
    if not path.is_absolute():
        path = base_dir / path
    return path if path.is_file() else None


def _strip_front_matter(text: str) -> tuple[str, str]:
    """去掉开头的 YAML front matter，返回 (正文, 标题)。"""
    title = ""
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if m:
        t = re.search(r"^title:\s*[\"']?(.*?)[\"']?\s*$", m.group(1), re.M)
        title = t.group(1) if t else ""
        text = text[m.end():]
    return text, title


def _plain(node) -> str:
    if node.type in ("text", "code_inline", "html_inline"):
        return node.content
    if node.type in ("softbreak", "hardbreak"):
        return " "
    if node.type == "image":
        return node.attrs.get("alt", "") or "".join(_plain(c) for c in node.children)
    return "".join(_plain(c) for c in node.children)


# ── HTML ─────────────────────────────────────────────────────────────────
_CSS = """
:root { color-scheme: light dark; --fg:#1f2328; --muted:#59636e; --line:#d1d9e0; --code:#f6f8fa; --link:#0969da; }
@media (prefers-color-scheme: dark) { :root { --fg:#e6edf3; --muted:#9198a1; --line:#3d444d; --code:#151b23; --link:#4493f8; } body { background:#0d1117; } }
body { margin:0; color:var(--fg); font:16px/1.7 -apple-system, "Segoe UI", "PingFang SC", "Microsoft YaHei", sans-serif; }
main { max-width:860px; margin:0 auto; padding:40px 24px 80px; }
h1,h2,h3,h4,h5,h6 { line-height:1.3; margin:1.6em 0 .6em; }
h1 { font-size:2em; padding-bottom:.3em; border-bottom:1px solid var(--line); }
h2 { font-size:1.5em; padding-bottom:.3em; border-bottom:1px solid var(--line); }
a { color:var(--link); text-decoration:none; } a:hover { text-decoration:underline; }
code { font-family:Consolas, Menlo, monospace; font-size:.9em; background:var(--code); padding:.15em .35em; border-radius:4px; }
pre { background:var(--code); padding:14px 16px; border-radius:8px; overflow:auto; line-height:1.5; }
pre code { background:none; padding:0; }
blockquote { margin:0; padding:0 1em; color:var(--muted); border-left:4px solid var(--line); }
table { border-collapse:collapse; display:block; overflow:auto; }
th,td { border:1px solid var(--line); padding:6px 13px; } th { background:var(--code); }
img { max-width:100%; }
hr { border:0; border-top:1px solid var(--line); margin:2em 0; }
li.task { list-style:none; } li.task input { margin:0 .5em 0 -1.4em; }
""".strip()


def _inline_images_html(html: str, base_dir: Path) -> str:
    def repl(m: re.Match) -> str:
        src = html_lib.unescape(m.group(2))
        path = _local_image(src, base_dir)
        if path is None:
            return m.group(0)
        mime = mimetypes.guess_type(path.name)[0] or "image/png"
        data = base64.b64encode(path.read_bytes()).decode("ascii")
        return f'{m.group(1)}data:{mime};base64,{data}{m.group(3)}'
    return re.sub(r'(<img[^>]*?\ssrc=")([^"]+)(")', repl, html)


def md_to_html_string(text: str, base_dir: Path, title: str = "") -> str:
    text, fm_title = _strip_front_matter(text)
    body = _parser().render(text)
    body = re.sub(r"<li>(\s*(?:<p>)?)\[([ xX])\]\s+",
                  lambda m: f'<li class="task">{m.group(1)}<input type="checkbox" disabled'
                            f'{" checked" if m.group(2) != " " else ""}> ', body)
    body = _inline_images_html(body, base_dir)
    if not title:
        m = re.search(r"<h1>(.*?)</h1>", body, re.S)
        title = re.sub(r"<[^>]+>", "", m.group(1)) if m else fm_title
    return (f'<!DOCTYPE html>\n<html lang="zh-CN">\n<head>\n<meta charset="utf-8">\n'
            f'<meta name="viewport" content="width=device-width, initial-scale=1">\n'
            f"<title>{html_lib.escape(title or fm_title)}</title>\n<style>\n{_CSS}\n</style>\n</head>\n"
            f"<body>\n<main>\n{body}</main>\n</body>\n</html>\n")


# ── Word ─────────────────────────────────────────────────────────────────
class _DocxWriter:
    def __init__(self, base_dir: Path) -> None:
        import docx
        from docx.shared import Pt

        self.base_dir = base_dir
        self.doc = docx.Document()
        normal = self.doc.styles["Normal"]
        normal.font.size = Pt(11)
        self._set_east_asia(normal.element, "Microsoft YaHei")
        from docx.enum.style import WD_STYLE_TYPE
        code = self.doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
        code.font.name = _MONO
        code.font.size = Pt(9.5)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(4)
        self._shade(code.element.get_or_add_pPr(), "F2F2F2")
        self.list_stack: list[str] = []   # "bullet" / "ordered"
        self.quote = 0

    @staticmethod
    def _set_east_asia(el, font: str) -> None:
        from docx.oxml.ns import qn

        rpr = el.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fonts)
        fonts.set(qn("w:eastAsia"), font)

    # 段落样式
    def _style(self, name: str) -> str | None:
        try:
            self.doc.styles[name]
            return name
        except KeyError:
            return None

    def _para(self, style: str | None = None):
        p = self.doc.add_paragraph(style=self._style(style) if style else None)
        if self.quote:
            from docx.shared import Pt
            p.paragraph_format.left_indent = Pt(18 * self.quote)
        return p

    # 行内
    def inline(self, p, nodes, fmt: dict | None = None) -> None:
        from docx.shared import RGBColor

        fmt = fmt or {}
        for n in nodes:
            t = n.type
            if t == "text":
                self._run(p, n.content, fmt)
            elif t == "softbreak":
                self._run(p, " ", fmt)
            elif t == "hardbreak":
                p.add_run().add_break()
            elif t == "code_inline":
                self._run(p, n.content, {**fmt, "code": True})
            elif t == "strong":
                self.inline(p, n.children, {**fmt, "bold": True})
            elif t == "em":
                self.inline(p, n.children, {**fmt, "italic": True})
            elif t == "s":
                self.inline(p, n.children, {**fmt, "strike": True})
            elif t == "link":
                self._hyperlink(p, n.attrs.get("href", ""), _plain(n), fmt)
            elif t == "image":
                path = _local_image(str(n.attrs.get("src", "")), self.base_dir)
                if path is not None:
                    try:
                        self._picture(p, path)
                        continue
                    except Exception:
                        pass
                alt = _plain(n) or str(n.attrs.get("src", ""))
                r = self._run(p, f"[图片：{alt}]", fmt)
                r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            elif t == "html_inline":
                if not re.fullmatch(r"<br\s*/?>", n.content.strip(), re.I):
                    continue
                p.add_run().add_break()
            else:
                self.inline(p, n.children, fmt)

    def _run(self, p, text: str, fmt: dict):
        from docx.shared import RGBColor

        r = p.add_run(text)
        if fmt.get("bold"):
            r.bold = True
        if fmt.get("italic"):
            r.italic = True
        if fmt.get("strike"):
            r.font.strike = True
        if fmt.get("code"):
            r.font.name = _MONO
            self._shade(r._r.get_or_add_rPr(), "F2F2F2")
        if fmt.get("link"):
            r.font.underline = True
            r.font.color.rgb = RGBColor(0x09, 0x69, 0xDA)
        return r

    @staticmethod
    def _shade(parent, fill: str) -> None:
        from docx.oxml.ns import qn

        shd = parent.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill})
        parent.append(shd)

    def _hyperlink(self, p, href: str, text: str, fmt: dict) -> None:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml.ns import qn

        if not href:
            self._run(p, text, fmt)
            return
        if href.startswith("#"):
            self._run(p, text, {**fmt, "link": True})
            return
        r_id = p.part.relate_to(href, RT.HYPERLINK, is_external=True)
        link = p._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
        run = self._run(p, text, {**fmt, "link": True})
        link.append(run._r)
        p._p.append(link)

    def _picture(self, p, path: Path) -> None:
        from docx.shared import Inches
        from PIL import Image

        with Image.open(path) as im:
            w = im.size[0]
            dpi = im.info.get("dpi", (96, 96))[0] or 96
            if im.format not in ("PNG", "JPEG", "GIF", "BMP", "TIFF"):
                buf = io.BytesIO()
                im.convert("RGBA").save(buf, "PNG")
                buf.seek(0)
                source: object = buf
            else:
                source = str(path)
        width = min(6.0, w / float(dpi))
        p.add_run().add_picture(source, width=Inches(max(0.2, width)))

    # 块
    def block(self, node) -> None:
        check_cancelled()
        t = node.type
        if t == "heading":
            level = int(node.tag[1])
            p = self.doc.add_heading(level=level)
            self.inline(p, node.children[0].children if node.children else [])
        elif t == "paragraph":
            self._paragraph(node)
        elif t in ("bullet_list", "ordered_list"):
            self.list_stack.append("ordered" if t == "ordered_list" else "bullet")
            for item in node.children:
                self._list_item(item)
            self.list_stack.pop()
        elif t in ("fence", "code_block"):
            self._code(node.content.rstrip("\n"))
        elif t == "blockquote":
            self.quote += 1
            for c in node.children:
                self.block(c)
            self.quote -= 1
        elif t == "table":
            self._table(node)
        elif t == "hr":
            self._hr()
        elif t == "html_block":
            text = re.sub(r"<[^>]+>", "", node.content).strip()
            if text:
                self._para().add_run(text)
        else:
            for c in node.children:
                self.block(c)

    def _paragraph(self, node, style: str | None = None) -> None:
        p = self._para(style or ("Quote" if self.quote else None))
        inline = node.children[0].children if node.children else []
        self.inline(p, inline)

    def _list_item(self, item) -> None:
        depth = len(self.list_stack)
        kind = self.list_stack[-1]
        base = "List Number" if kind == "ordered" else "List Bullet"
        style = base if depth == 1 else f"{base} {min(depth, 3)}"
        first = True
        for c in item.children:
            if c.type == "paragraph" and first:
                p = self._para(style)
                inline = list(c.children[0].children) if c.children else []
                if inline and inline[0].type == "text" and _TASK_RE.match(inline[0].content):
                    m = _TASK_RE.match(inline[0].content)
                    assert m is not None
                    p.add_run("☑ " if m.group(1) != " " else "☐ ")
                    self._run(p, inline[0].content[m.end():], {})
                    inline = inline[1:]
                self.inline(p, inline)
                first = False
            else:
                self.block(c)

    def _code(self, code: str) -> None:
        from docx.oxml.ns import qn

        p = self._para("Code")
        for i, line in enumerate(code.split("\n")):
            if i:
                p.add_run().add_break()
            r = p.add_run(line)
            r.font.name = _MONO
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")

    def _table(self, node) -> None:
        rows = []
        for section in node.children:  # thead / tbody
            for tr in section.children:
                rows.append([cell for cell in tr.children])
        if not rows:
            return
        cols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=cols)
        if self._style("Table Grid"):
            table.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                p = table.cell(i, j).paragraphs[0]
                inline = cell.children[0].children if cell.children else []
                self.inline(p, inline, {"bold": True} if i == 0 else {})
        self.doc.add_paragraph()

    def _hr(self) -> None:
        from docx.oxml.ns import qn

        p = self._para()
        ppr = p._p.get_or_add_pPr()
        bdr = ppr.makeelement(qn("w:pBdr"), {})
        bottom = bdr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "6",
                                                  qn("w:space"): "1", qn("w:color"): "BFBFBF"})
        bdr.append(bottom)
        ppr.append(bdr)


def md_to_docx_file(text: str, base_dir: Path, out_path: Path) -> None:
    text, title = _strip_front_matter(text)
    w = _DocxWriter(base_dir)
    if title:
        w.doc.core_properties.title = title
    for node in _tree(text).children:
        w.block(node)
    w.doc.save(str(out_path))


# ── PDF ──────────────────────────────────────────────────────────────────
class _PdfWriter:
    def __init__(self, base_dir: Path, sample_text: str) -> None:
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics

        from core.fonts import reportlab_font

        self.base_dir = base_dir
        chars = "".join(sorted({ch for ch in sample_text if not ch.isspace() and ord(ch) > 127}))[:3000]
        self.font = reportlab_font(chars or "A")
        if self.font != "Helvetica":
            pdfmetrics.registerFontFamily(self.font, normal=self.font, bold=self.font, italic=self.font,
                                          boldItalic=self.font)
        self.mono = "Courier" if sample_text.isascii() else self.font
        self.colors = colors
        base = ParagraphStyle("body", fontName=self.font, fontSize=10.5, leading=17, spaceAfter=6,
                              wordWrap="CJK")
        self.styles = {"body": base}
        for lvl, size in enumerate((22, 18, 15, 13, 12, 11), start=1):
            self.styles[f"h{lvl}"] = ParagraphStyle(f"h{lvl}", parent=base, fontSize=size, leading=size * 1.35,
                                                    spaceBefore=size * 0.8, spaceAfter=size * 0.4)
        self.styles["code"] = ParagraphStyle("code", parent=base, fontName=self.mono, fontSize=9, leading=13,
                                             spaceAfter=0)
        self.styles["quote"] = ParagraphStyle("quote", parent=base, textColor=colors.HexColor("#59636E"),
                                              leftIndent=14)
        self.styles["cell"] = ParagraphStyle("cell", parent=base, fontSize=9.5, leading=14, spaceAfter=0)
        self.story: list = []
        self.max_width = 0.0

    def _esc(self, s: str) -> str:
        return html_lib.escape(s, quote=False)

    def inline(self, nodes) -> str:
        out: list[str] = []
        for n in nodes:
            t = n.type
            if t == "text":
                out.append(self._esc(n.content))
            elif t == "softbreak":
                out.append(" ")
            elif t == "hardbreak":
                out.append("<br/>")
            elif t == "code_inline":
                out.append(f'<font face="{self.mono}" backColor="#F3F4F6">{self._esc(n.content)}</font>')
            elif t == "strong":
                out.append(f"<b>{self.inline(n.children)}</b>")
            elif t == "em":
                out.append(f"<i>{self.inline(n.children)}</i>")
            elif t == "s":
                out.append(f"<strike>{self.inline(n.children)}</strike>")
            elif t == "link":
                href = html_lib.escape(str(n.attrs.get("href", "")))
                out.append(f'<a href="{href}" color="#0969DA">{self.inline(n.children)}</a>')
            elif t == "image":
                out.append(self._esc(f"[图片：{_plain(n)}]"))
            elif t == "html_inline":
                if re.fullmatch(r"<br\s*/?>", n.content.strip(), re.I):
                    out.append("<br/>")
            else:
                out.append(self.inline(n.children))
        return "".join(out)

    def _images_in(self, nodes) -> list[Path]:
        found = []
        for n in nodes:
            if n.type == "image":
                p = _local_image(str(n.attrs.get("src", "")), self.base_dir)
                if p is not None:
                    found.append(p)
            else:
                found.extend(self._images_in(n.children))
        return found

    def _image(self, path: Path):
        from reportlab.platypus import Image as RLImage

        img = RLImage(str(path))
        w, h = img.imageWidth, img.imageHeight
        scale = min(1.0, self.max_width / w if w else 1.0)
        img.drawWidth, img.drawHeight = w * scale * 0.75, h * scale * 0.75
        if img.drawWidth > self.max_width:
            k = self.max_width / img.drawWidth
            img.drawWidth, img.drawHeight = img.drawWidth * k, img.drawHeight * k
        return img

    def blocks(self, nodes, quote: bool = False) -> list:
        from reportlab.platypus import (
            HRFlowable,
            ListFlowable,
            ListItem,
            Paragraph,
            Preformatted,
            Spacer,
            Table,
            TableStyle,
        )

        flow: list = []
        for node in nodes:
            check_cancelled()
            t = node.type
            if t == "heading":
                inline = node.children[0].children if node.children else []
                flow.append(Paragraph(self.inline(inline), self.styles[node.tag]))
            elif t == "paragraph":
                inline = node.children[0].children if node.children else []
                images = [n for n in inline if n.type == "image"]
                if images and len(images) == len([n for n in inline if n.type not in ("softbreak",)]):
                    for p in self._images_in(inline):
                        try:
                            flow.append(self._image(p))
                        except Exception:
                            pass
                    continue
                text = self.inline(inline)
                m = _TASK_RE.match(text)
                if m:
                    text = ("[x] " if m.group(1) != " " else "[ ] ") + text[m.end():]
                flow.append(Paragraph(text, self.styles["quote" if quote else "body"]))
            elif t in ("bullet_list", "ordered_list"):
                items = [ListItem(self.blocks(item.children, quote), leftIndent=18)
                         for item in node.children]
                if t == "ordered_list":
                    flow.append(ListFlowable(items, bulletType="1", bulletFontName=self.font,
                                             bulletFontSize=10, leftIndent=18,
                                             start=int(node.attrs.get("start", 1) or 1)))
                else:
                    flow.append(ListFlowable(items, bulletType="bullet", start="•", bulletFontName="Helvetica",
                                             bulletFontSize=8, leftIndent=18))
            elif t in ("fence", "code_block"):
                code = Table([[Preformatted(node.content.rstrip("\n"), self.styles["code"])]],
                             colWidths=[self.max_width])
                code.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), self.colors.HexColor("#F3F4F6")),
                    ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                    ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ]))
                flow.append(code)
                flow.append(Spacer(1, 8))
            elif t == "blockquote":
                flow.extend(self.blocks(node.children, True))
            elif t == "table":
                rows = []
                for section in node.children:
                    for tr in section.children:
                        rows.append([Paragraph(self.inline(c.children[0].children if c.children else []),
                                               self.styles["cell"]) for c in tr.children])
                if rows:
                    cols = max(len(r) for r in rows)
                    rows = [r + [""] * (cols - len(r)) for r in rows]
                    tbl = Table(rows, colWidths=[self.max_width / cols] * cols, repeatRows=1)
                    tbl.setStyle(TableStyle([
                        ("GRID", (0, 0), (-1, -1), 0.5, self.colors.HexColor("#D0D7DE")),
                        ("BACKGROUND", (0, 0), (-1, 0), self.colors.HexColor("#F3F4F6")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]))
                    flow.append(tbl)
                    flow.append(Spacer(1, 8))
            elif t == "hr":
                flow.append(HRFlowable(width="100%", thickness=0.6, color=self.colors.HexColor("#D0D7DE"),
                                       spaceBefore=6, spaceAfter=10))
            elif t == "html_block":
                text = re.sub(r"<[^>]+>", "", node.content).strip()
                if text:
                    flow.append(Paragraph(self._esc(text), self.styles["body"]))
            else:
                flow.extend(self.blocks(node.children, quote))
        return flow


def md_to_pdf_file(text: str, base_dir: Path, out_path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate

    text, title = _strip_front_matter(text)
    w = _PdfWriter(base_dir, text)
    doc = SimpleDocTemplate(str(out_path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm, title=title or out_path.stem)
    w.max_width = doc.width
    story = w.blocks(_tree(text).children)
    if not story:
        raise RuntimeError("Markdown 文件没有内容")
    doc.build(story)


def from_markdown(
    input_file: Path,
    output_dir: Path,
    target: str = "docx",
    progress_callback: ProgressCallback | None = None,
    **_,
) -> TaskResult:
    """一个 .md 文件转成 Word / HTML / PDF。"""
    t0 = time.time()
    try:
        if target not in TARGETS:
            raise ValueError(f"不支持的目标格式：{target}")
        output_dir.mkdir(parents=True, exist_ok=True)
        if progress_callback:
            progress_callback(0, 1, f"正在生成 {target.upper()}…")
        text = read_text_file(input_file).replace("\r\n", "\n")
        out_path = unique_path(output_dir / f"{input_file.stem}.{target}", [input_file])
        base_dir = input_file.parent
        if target == "html":
            out_path.write_text(md_to_html_string(text, base_dir), encoding="utf-8", newline="\n")
        elif target == "docx":
            md_to_docx_file(text, base_dir, out_path)
        else:
            md_to_pdf_file(text, base_dir, out_path)
        if progress_callback:
            progress_callback(1, 1, "转换完成")
        return TaskResult(status=TaskStatus.SUCCESS, output_files=[out_path], output_dir=output_dir,
                          duration_seconds=time.time() - t0)
    except TaskCancelled:
        return TaskResult(status=TaskStatus.CANCELLED, error_message="已取消", duration_seconds=time.time() - t0)
    except Exception as exc:
        return TaskResult(status=TaskStatus.FAILED, error_message=str(exc), duration_seconds=time.time() - t0)
