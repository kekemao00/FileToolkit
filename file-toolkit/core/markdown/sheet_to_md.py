"""表格文件转 Markdown：Excel（.xlsx）每个工作表一个表格，CSV 一个表格。"""
from __future__ import annotations

import csv
import datetime as dt
import io
from pathlib import Path

from core.markdown._common import MdDoc, read_text_file, table_md
from core.task_control import check_cancelled


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:.10g}"
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if v.time() == dt.time() else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return str(v)


def xlsx_to_md(path: Path) -> MdDoc:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[str] = []
    sheets = 0
    try:
        for ws in wb.worksheets:
            check_cancelled()
            rows = [[_fmt(v) for v in row] for row in ws.iter_rows(values_only=True)]
            md = table_md(rows)
            if not md:
                continue
            sheets += 1
            blocks.append(f"## {ws.title}\n\n{md}")
    finally:
        wb.close()
    # 只有一个工作表时不需要工作表标题
    if sheets == 1:
        blocks = [blocks[0].split("\n\n", 1)[1]]
    return MdDoc("\n\n".join(blocks), {"sheets": sheets})


def csv_to_md(path: Path) -> MdDoc:
    text = read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return MdDoc(table_md(rows), {"rows": max(0, len(rows) - 1)})
