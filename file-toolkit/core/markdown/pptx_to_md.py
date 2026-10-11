"""PowerPoint（.pptx）转 Markdown：每页一个二级标题，正文保留层级，表格、图片、备注。"""
from __future__ import annotations

from pathlib import Path

from core.markdown._common import ImageSink, MdDoc, table_md
from core.task_control import check_cancelled


def _shapes(shapes):
    """按从上到下、从左到右的阅读顺序展开（组合形状递归展开）。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    flat = []
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            flat.extend(_shapes(sh.shapes))
        else:
            flat.append(sh)
    return sorted(flat, key=lambda s: ((s.top or 0) // 20000, s.left or 0))


def _text_frame(tf) -> str:
    paras = [(p.level, "".join(r.text for r in p.runs).strip() or p.text.strip()) for p in tf.paragraphs]
    paras = [(lvl, t) for lvl, t in paras if t]
    if not paras:
        return ""
    if len(paras) == 1 and paras[0][0] == 0:
        return paras[0][1]
    return "\n".join("    " * lvl + "- " + t.replace("\n", " ") for lvl, t in paras)


def pptx_to_md(path: Path, images: ImageSink) -> MdDoc:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    blocks: list[str] = []
    for n, slide in enumerate(prs.slides, start=1):
        check_cancelled()
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip().replace("\n", " ") if title_shape is not None \
            and title_shape.has_text_frame else ""
        blocks.append(f"## 第 {n} 页" + (f"：{title}" if title else ""))
        for sh in _shapes(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[cell.text for cell in row.cells] for row in sh.table.rows]
                md = table_md(rows)
                if md:
                    blocks.append(md)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE and images.enabled:
                try:
                    md = images.add(sh.image.blob, sh.image.ext, sh.name if not sh.name.startswith("Picture") else "")
                except (AttributeError, ValueError, KeyError):
                    md = ""
                if md:
                    blocks.append(md)
            elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                md = _text_frame(sh.text_frame)
                if md:
                    blocks.append(md)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                blocks.append("> 备注：" + notes.replace("\n", "\n> "))
    meta: dict[str, str | int] = {"slides": len(prs.slides)}
    if (prs.core_properties.title or "").strip():
        meta["title"] = prs.core_properties.title.strip()
    return MdDoc("\n\n".join(blocks), meta)
