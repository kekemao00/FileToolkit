"""文档 ↔ Markdown 转换。"""
from pathlib import Path

import docx
import pytest
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as PptInches
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from core.markdown._common import (
    IMAGE_EMBED,
    IMAGE_SKIP,
    ImageSink,
    estimate_tokens,
    join_lines,
    normalize_md,
    table_md,
)
from core.markdown.converter import markdown_stats, merge_to_markdown, to_markdown
from core.markdown.from_md import from_markdown, md_to_html_string
from core.models import TaskStatus


def _png(path: Path) -> Path:
    Image.new("RGB", (60, 40), (10, 120, 200)).save(path)
    return path


def _docx(path: Path, image: Path | None = None) -> Path:
    d = docx.Document()
    d.add_heading("需求说明", 1)
    p = d.add_paragraph("这是")
    p.add_run("加粗").bold = True
    p.add_run("和")
    p.add_run("斜体").italic = True
    d.add_paragraph("第一项", style="List Bullet")
    d.add_paragraph("子项", style="List Bullet 2")
    d.add_paragraph("步骤一", style="List Number")
    d.add_paragraph("步骤二", style="List Number")
    d.add_heading("数据", 2)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "名称", "数量"
    t.cell(1, 0).text, t.cell(1, 1).text = "苹果", "3|4"
    if image:
        d.add_picture(str(image), width=Inches(1))
    d.save(path)
    return path


def _pdf(path: Path, pages: int = 3) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    for n in range(pages):
        c.setFont("Helvetica", 8)
        c.drawString(250, 815, "ACME Internal Report")
        c.setFont("Helvetica-Bold", 20)
        c.drawString(72, 760, f"Chapter {n + 1} Overview")
        c.setFont("Helvetica", 11)
        y = 730
        for line in ["This is the first line of a paragraph that has been wrapped by the",
                     "layout engine and should be joined back into a single paragraph.",
                     "A second paragraph starts after a gap and ends with a full stop here."]:
            c.drawString(72, y, line)
            y -= 14 if "layout" in line or "first" in line else 0
            if "layout" in line:
                y -= 16
        c.drawString(72, 650, "• bullet item one")
        c.drawString(72, 636, "• bullet item two")
        c.setFont("Helvetica", 9)
        c.drawString(290, 30, str(n + 1))
        c.showPage()
    c.save()
    return path


# ── 公共工具 ─────────────────────────────────────────────────────────────
def test_table_md_escapes_pipes_and_drops_empty_columns() -> None:
    md = table_md([["a", "", "b"], ["1", "", "x|y"], ["", "", ""]])
    assert md.splitlines() == ["| a | b |", "| --- | --- |", "| 1 | x\\|y |"]


def test_normalize_md_collapses_blank_lines_but_keeps_code() -> None:
    text = "a  \n\n\n\nb\n```\nx\n\n\n\ny\n```"
    assert normalize_md(text) == "a\n\nb\n```\nx\n\n\n\ny\n```\n"


def test_join_lines_handles_cjk_and_hyphenation() -> None:
    assert join_lines("中文第一行", "第二行") == "中文第一行第二行"
    assert join_lines("hyphen-", "ated word") == "hyphenated word"
    assert join_lines("hello", "world") == "hello world"


def test_estimate_tokens_counts_cjk_by_char() -> None:
    assert estimate_tokens("你好世界") == 4
    assert 4 <= estimate_tokens("hello brave new world") <= 6


def test_image_sink_modes(tmp_path: Path) -> None:
    data = _png(tmp_path / "a.png").read_bytes()
    assert ImageSink(IMAGE_SKIP).add(data, "png") == ""
    assert ImageSink(IMAGE_EMBED).add(data, "png").startswith("![](data:image/png;base64,")
    sink = ImageSink("extract", tmp_path / "doc_assets", tmp_path)
    first = sink.add(data, "png")
    assert first == "![](doc_assets/image_001.png)"
    assert sink.add(data, "png") == first  # 同一张图只保存一次
    assert len(sink.files) == 1


# ── 转 Markdown ──────────────────────────────────────────────────────────
def test_docx_to_markdown_keeps_structure_and_images(tmp_path: Path) -> None:
    src = _docx(tmp_path / "需求.docx", _png(tmp_path / "pic.png"))
    result = to_markdown(src, tmp_path / "out")
    assert result.status == TaskStatus.SUCCESS, result.error_message
    md = result.output_files[0].read_text(encoding="utf-8")
    assert "# 需求说明" in md
    assert "## 数据" in md
    assert "这是**加粗**和*斜体*" in md
    assert "- 第一项\n    - 子项" in md
    assert "1. 步骤一\n2. 步骤二" in md
    assert "| 名称 | 数量 |" in md and "| 苹果 | 3\\|4 |" in md
    assert "![](需求_assets/image_001.png)" in md
    assert (tmp_path / "out" / "需求_assets" / "image_001.png").is_file()


def test_pdf_to_markdown_detects_headings_paragraphs_and_strips_page_furniture(tmp_path: Path) -> None:
    result = to_markdown(_pdf(tmp_path / "r.pdf"), tmp_path / "out", page_markers=True)
    assert result.status == TaskStatus.SUCCESS, result.error_message
    md = result.output_files[0].read_text(encoding="utf-8")
    assert "# Chapter 1 Overview" in md and "# Chapter 3 Overview" in md
    assert ("This is the first line of a paragraph that has been wrapped by the layout engine "
            "and should be joined back into a single paragraph.") in md
    assert "- bullet item one" in md and "- bullet item two" in md
    assert "<!-- 第 2 页 -->" in md
    assert "ACME Internal Report" not in md  # 页眉
    assert "\n1\n" not in md  # 页码


def test_scanned_pdf_without_tesseract_fails_with_hint(tmp_path: Path, monkeypatch) -> None:
    c = canvas.Canvas(str(tmp_path / "scan.pdf"), pagesize=A4)
    c.showPage()
    c.save()
    monkeypatch.setattr("core.ocr.client.find_tesseract", lambda: None)
    result = to_markdown(tmp_path / "scan.pdf", tmp_path / "out")
    assert result.status == TaskStatus.FAILED
    assert "Tesseract" in (result.error_message or "")


def test_pptx_to_markdown_has_one_section_per_slide(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "季度总结"
    tf = slide.placeholders[1].text_frame
    tf.text = "营收增长"
    sub = tf.add_paragraph()
    sub.text, sub.level = "同比 +20%", 1
    slide.shapes.add_picture(str(_png(tmp_path / "p.png")), PptInches(1), PptInches(5))
    slide.notes_slide.notes_text_frame.text = "讲 3 分钟"
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "谢谢"
    prs.save(tmp_path / "deck.pptx")

    result = to_markdown(tmp_path / "deck.pptx", tmp_path / "out", image_mode=IMAGE_SKIP)
    md = result.output_files[0].read_text(encoding="utf-8")
    assert "## 第 1 页：季度总结" in md and "## 第 2 页：谢谢" in md
    assert "- 营收增长\n    - 同比 +20%" in md
    assert "> 备注：讲 3 分钟" in md
    assert "![" not in md
    assert not (tmp_path / "out" / "deck_assets").exists()


def test_xlsx_and_csv_to_markdown_tables(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["月份", "金额"])
    ws.append(["1月", 100.0])
    ws.append(["2月", 120.5])
    wb.create_sheet("成本").append(["项目", "费用"])
    wb["成本"].append(["房租", 50])
    wb.save(tmp_path / "data.xlsx")
    md = to_markdown(tmp_path / "data.xlsx", tmp_path / "out").output_files[0].read_text(encoding="utf-8")
    assert "## 销售" in md and "## 成本" in md
    assert "| 1月 | 100 |" in md and "| 2月 | 120.5 |" in md

    (tmp_path / "list.csv").write_bytes("姓名,城市\n张三,北京\n".encode("gbk"))
    md = to_markdown(tmp_path / "list.csv", tmp_path / "out").output_files[0].read_text(encoding="utf-8")
    assert md.strip() == "| 姓名 | 城市 |\n| --- | --- |\n| 张三 | 北京 |"


def test_html_to_markdown_drops_noise(tmp_path: Path) -> None:
    _png(tmp_path / "pic.png")
    (tmp_path / "page.html").write_text(
        "<html><head><title>T</title><script>var x=1</script></head><body><nav>菜单</nav>"
        "<article><h1>标题</h1><p>Hello <b>world</b> <a href='https://e.com'>link</a></p>"
        "<ul><li>a<ul><li>b</li></ul></li><li>c</li></ul>"
        "<pre><code class='language-py'>print(1)\n</code></pre>"
        "<table><tr><th>x</th><th>y</th></tr><tr><td>1</td><td>2</td></tr></table>"
        "<img src='pic.png' alt='图'></article></body></html>", encoding="utf-8")
    md = to_markdown(tmp_path / "page.html", tmp_path / "out").output_files[0].read_text(encoding="utf-8")
    assert md.startswith("# 标题\n")
    assert "Hello **world** [link](https://e.com)" in md
    assert "- a\n    - b\n- c" in md
    assert "```py\nprint(1)\n```" in md
    assert "| x | y |" in md
    assert "![图](page_assets/image_001.png)" in md
    assert "菜单" not in md and "var x" not in md


def test_front_matter_and_unsupported_format(tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_text("第一行\r\n\r\n\r\n第二行", encoding="utf-8")
    md = to_markdown(tmp_path / "a.txt", tmp_path / "out", include_meta=True).output_files[0].read_text(
        encoding="utf-8")
    assert md.startswith('---\nsource: "a.txt"\nformat: "文本"\nconverted: ')
    assert md.endswith("第一行\n\n第二行\n")

    (tmp_path / "x.bin").write_bytes(b"\0")
    assert to_markdown(tmp_path / "x.bin", tmp_path / "out").status == TaskStatus.FAILED


def test_merge_to_one_markdown_demotes_headings(tmp_path: Path) -> None:
    a = _docx(tmp_path / "a.docx")
    (tmp_path / "b.txt").write_text("# 已有标题\n正文", encoding="utf-8")
    (tmp_path / "bad.pdf").write_bytes(b"not a pdf")
    result = merge_to_markdown([a, tmp_path / "b.txt", tmp_path / "bad.pdf"], tmp_path / "out" / "合并.md",
                               image_mode=IMAGE_SKIP)
    assert result.status == TaskStatus.SUCCESS
    assert len(result.warnings) == 1 and result.warnings[0].startswith("bad.pdf")
    md = result.output_files[0].read_text(encoding="utf-8")
    assert "# a.docx\n\n## 需求说明" in md
    assert "# b.txt\n\n## 已有标题" in md
    chars, tokens = markdown_stats(result.output_files)
    assert chars == len(md) and tokens > 0


@pytest.mark.parametrize("ext", ["doc", "rtf"])
def test_legacy_formats_need_libreoffice(tmp_path: Path, monkeypatch, ext: str) -> None:
    monkeypatch.setattr("core.markdown.converter.get_libreoffice_path", lambda: None)
    (tmp_path / f"old.{ext}").write_bytes(b"x")
    result = to_markdown(tmp_path / f"old.{ext}", tmp_path / "out")
    assert result.status == TaskStatus.FAILED
    assert "LibreOffice" in (result.error_message or "")


# ── Markdown 转出 ────────────────────────────────────────────────────────
_MD = """---
title: "周报"
---
# 本周周报

这是**加粗**、*斜体*、~~删除~~和 `code`，还有[链接](https://example.com)。

## 进展

- 完成 A
    - 子项 a1
- [x] 已完成任务

1. 第一步
2. 第二步

> 引用

```python
def hi():
    print("你好")
```

| 指标 | 数值 |
| --- | --- |
| 营收 | 100 |

---

![图](pic.png)
"""


def test_markdown_to_docx_round_trips(tmp_path: Path) -> None:
    _png(tmp_path / "pic.png")
    (tmp_path / "周报.md").write_text(_MD, encoding="utf-8")
    result = from_markdown(tmp_path / "周报.md", tmp_path / "out", "docx")
    assert result.status == TaskStatus.SUCCESS, result.error_message
    d = docx.Document(str(result.output_files[0]))
    assert d.core_properties.title == "周报"
    styles = [p.style.name for p in d.paragraphs if p.text.strip()]
    assert styles[0] == "Heading 1" and "Heading 2" in styles and "List Bullet 2" in styles
    assert d.tables[0].cell(1, 0).text == "营收"
    assert d.inline_shapes and len(d.inline_shapes) == 1

    back = to_markdown(result.output_files[0], tmp_path / "back", image_mode=IMAGE_SKIP)
    md = back.output_files[0].read_text(encoding="utf-8")
    assert "# 本周周报" in md and "**加粗**" in md and "~~删除~~" in md and "[链接](https://example.com)" in md
    assert "1. 第一步\n2. 第二步" in md
    assert '```\ndef hi():\n    print("你好")\n```' in md
    assert "| 指标 | 数值 |" in md


def test_markdown_to_html_is_self_contained(tmp_path: Path) -> None:
    _png(tmp_path / "pic.png")
    html = md_to_html_string(_MD, tmp_path)
    assert "<title>本周周报</title>" in html
    assert "<table>" in html and "<s>删除</s>" in html
    assert 'type="checkbox" disabled checked' in html
    assert 'src="data:image/png;base64,' in html
    assert "title: " not in html  # front matter 不出现在正文


def test_markdown_to_pdf(tmp_path: Path) -> None:
    import pdfplumber

    (tmp_path / "a.md").write_text(_MD, encoding="utf-8")
    result = from_markdown(tmp_path / "a.md", tmp_path / "out", "pdf")
    assert result.status == TaskStatus.SUCCESS, result.error_message
    with pdfplumber.open(str(result.output_files[0])) as pdf:
        text = "".join(p.extract_text() or "" for p in pdf.pages)
    assert "第一步" in text or "1" in text
    assert "def hi():" in text


def test_from_markdown_rejects_unknown_target(tmp_path: Path) -> None:
    (tmp_path / "a.md").write_text("# x", encoding="utf-8")
    assert from_markdown(tmp_path / "a.md", tmp_path / "out", "epub").status == TaskStatus.FAILED
